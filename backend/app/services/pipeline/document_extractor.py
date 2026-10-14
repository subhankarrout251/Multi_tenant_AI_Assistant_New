"""
app/services/pipeline/document_extractor.py
──────────────────────────────────────────────────────────────────────
Unified document extractor for PDF (.pdf) and Word (.docx) files.
Extracts clean, structured text with file name, section/page numbers,
and document type metadata.
──────────────────────────────────────────────────────────────────────
"""

import os
import io
import re
import logging
from typing import BinaryIO, Union
from pypdf import PdfReader
import docx

logger = logging.getLogger(__name__)


def clean_text(text: str) -> str:
    """Normalize whitespace and linebreaks in extracted document text."""
    if not text:
        return ""
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = text.encode("utf-8", "ignore").decode("utf-8")
    return text.strip()


def extract_from_pdf(source: Union[str, bytes, BinaryIO], filename: str = "document.pdf") -> list[dict]:
    """Extracts text page-by-page from a PDF."""
    if isinstance(source, bytes):
        stream = io.BytesIO(source)
    elif isinstance(source, str):
        if not os.path.exists(source):
            raise FileNotFoundError(f"PDF not found: {source}")
        stream = open(source, "rb")
    else:
        stream = source

    try:
        reader = PdfReader(stream)
        pages = []
        clean_name = os.path.basename(filename)

        for idx, page in enumerate(reader.pages, start=1):
            raw_text = page.extract_text() or ""
            cleaned = clean_text(raw_text)

            if len(cleaned) < 30:
                continue

            pages.append({
                "url": f"{clean_name}#page={idx}",
                "page_title": f"{clean_name} - Page {idx}",
                "page_type": "pdf_document",
                "doc_type": "pdf",
                "text": cleaned,
                "page_number": idx,
                "source_file": clean_name,
            })

        logger.info("Extracted %d pages from PDF '%s'", len(pages), clean_name)
        return pages
    finally:
        if isinstance(source, str) and hasattr(stream, "close"):
            stream.close()


def extract_from_docx(source: Union[str, bytes, BinaryIO], filename: str = "document.docx") -> list[dict]:
    """Extracts text section-by-section from a Word .docx document."""
    if isinstance(source, bytes):
        stream = io.BytesIO(source)
    elif isinstance(source, str):
        if not os.path.exists(source):
            raise FileNotFoundError(f"DOCX file not found: {source}")
        stream = open(source, "rb")
    else:
        stream = source

    try:
        doc = docx.Document(stream)
        clean_name = os.path.basename(filename)

        paragraphs_text = []
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                paragraphs_text.append(txt)

        # Also extract table text
        for table in doc.tables:
            table_lines = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    table_lines.append(" | ".join(row_cells))
            if table_lines:
                paragraphs_text.append("\n".join(table_lines))

        full_text = "\n\n".join(paragraphs_text)
        cleaned = clean_text(full_text)

        if not cleaned:
            logger.warning("No readable text found in DOCX '%s'", clean_name)
            return []

        # Split long docx into ~1500 character logical sections
        SECTION_SIZE = 1500
        sections = []
        start = 0
        sec_idx = 1

        while start < len(cleaned):
            end = start + SECTION_SIZE
            chunk_slice = cleaned[start:end]
            last_break = chunk_slice.rfind("\n\n")
            if last_break > SECTION_SIZE // 2 and end < len(cleaned):
                actual_end = start + last_break + 2
            else:
                actual_end = end

            sec_text = cleaned[start:actual_end].strip()
            if sec_text:
                sections.append({
                    "url": f"{clean_name}#section={sec_idx}",
                    "page_title": f"{clean_name} - Section {sec_idx}",
                    "page_type": "docx_document",
                    "doc_type": "docx",
                    "text": sec_text,
                    "page_number": sec_idx,
                    "source_file": clean_name,
                })
                sec_idx += 1

            start = actual_end

        logger.info("Extracted %d section(s) from DOCX '%s'", len(sections), clean_name)
        return sections
    finally:
        if isinstance(source, str) and hasattr(stream, "close"):
            stream.close()


def extract_document(source: Union[str, bytes, BinaryIO], filename: str) -> list[dict]:
    """Dispatches extraction according to file extension (.pdf or .docx)."""
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return extract_from_pdf(source, filename)
    elif ext == ".docx":
        return extract_from_docx(source, filename)
    else:
        raise ValueError(f"Unsupported document format '{ext}'. Only .pdf and .docx are supported.")


def extract_documents_from_directory(dir_path: str) -> list[dict]:
    """Extracts all PDF and DOCX files in a directory."""
    if not os.path.isdir(dir_path):
        raise NotADirectoryError(f"Directory not found: {dir_path}")

    all_pages = []
    valid_exts = (".pdf", ".docx")
    files = [f for f in os.listdir(dir_path) if f.lower().endswith(valid_exts)]

    if not files:
        logger.warning("No PDF or DOCX files found in %s", dir_path)
        return []

    for fname in sorted(files):
        fpath = os.path.join(dir_path, fname)
        try:
            pages = extract_document(fpath, fname)
            all_pages.extend(pages)
        except Exception as e:
            logger.error("Failed to extract %s: %s", fname, e)

    return all_pages
